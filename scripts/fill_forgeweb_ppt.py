from __future__ import annotations

from copy import deepcopy
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


BODY_FONT = "Times New Roman"
TITLE_FONT = "Century Schoolbook"
PURPLE = RGBColor(112, 48, 160)
BLUE = RGBColor(0, 112, 192)
BLACK = RGBColor(0, 0, 0)
GRAY = RGBColor(68, 68, 68)
RED = RGBColor(214, 38, 43)
LIGHT_RED = RGBColor(252, 238, 239)
LIGHT_GRAY = RGBColor(247, 247, 247)
MID_GRAY = RGBColor(184, 184, 184)


def clear_text_frame(shape, *, margin=0.08):
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.word_wrap = True
    return frame


def disable_automatic_bullet(paragraph):
    properties = paragraph._p.get_or_add_pPr()
    for child in list(properties):
        if child.tag.endswith(("buNone", "buChar", "buAutoNum", "buBlip")):
            properties.remove(child)
    properties.insert(0, OxmlElement("a:buNone"))


def add_paragraph(frame, text, *, size=20, bold=False, color=BLACK, space_after=5, align=PP_ALIGN.LEFT):
    paragraph = frame.paragraphs[0] if len(frame.paragraphs) == 1 and not frame.paragraphs[0].text else frame.add_paragraph()
    paragraph.alignment = align
    disable_automatic_bullet(paragraph)
    paragraph.space_after = Pt(space_after)
    paragraph.line_spacing = 1.0
    run = paragraph.add_run()
    run.text = text
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return paragraph


def add_bullet(frame, text, *, size=20, bold=False, color=BLACK, space_after=6):
    return add_paragraph(frame, f"• {text}", size=size, bold=bold, color=color, space_after=space_after)


def set_body(slide, lines, *, size=20, body_index=1):
    frame = clear_text_frame(slide.shapes[body_index])
    for item in lines:
        if isinstance(item, tuple):
            text, bold = item
        else:
            text, bold = item, False
        add_bullet(frame, text, size=size, bold=bold)
    return frame


def set_paragraph_body(slide, paragraphs, *, size=20, body_index=1):
    frame = clear_text_frame(slide.shapes[body_index], margin=0.06)
    for text in paragraphs:
        add_paragraph(
            frame,
            text,
            size=size,
            space_after=14,
            align=PP_ALIGN.JUSTIFY,
        )
    return frame


def set_title(shape, text):
    frame = clear_text_frame(shape, margin=0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.name = BODY_FONT
    run.font.bold = True
    run.font.color.rgb = BLACK


def add_cropped_picture(slide, image_path, left, top, width, height):
    image_path = Path(image_path)
    with Image.open(image_path) as image:
        source_ratio = image.width / image.height
    box_ratio = width / height
    picture = slide.shapes.add_picture(str(image_path), Inches(left), Inches(top), Inches(width), Inches(height))
    if source_ratio > box_ratio:
        visible_fraction = box_ratio / source_ratio
        crop = (1 - visible_fraction) / 2
        picture.crop_left = crop
        picture.crop_right = crop
    elif source_ratio < box_ratio:
        visible_fraction = source_ratio / box_ratio
        crop = (1 - visible_fraction) / 2
        picture.crop_top = crop
        picture.crop_bottom = crop
    return picture


def add_caption(slide, text, left, top, width):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.25))
    frame = clear_text_frame(box, margin=0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    run.font.name = BODY_FONT
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.color.rgb = GRAY


def copy_footer_placeholders(source_slide, target_slide):
    for shape in source_slide.shapes:
        if not shape.is_placeholder:
            continue
        if int(shape.placeholder_format.type) not in {13, 15}:
            continue
        target_slide.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")


def bring_footer_placeholders_to_front(slide):
    for shape in list(slide.shapes):
        if not shape.is_placeholder or int(shape.placeholder_format.type) not in {13, 15}:
            continue
        element = shape.element
        slide.shapes._spTree.remove(element)
        slide.shapes._spTree.insert_element_before(element, "p:extLst")


def add_visible_template_footer(slide):
    for shape in list(slide.shapes):
        if shape.is_placeholder and int(shape.placeholder_format.type) == 15:
            slide.shapes._spTree.remove(shape.element)

    sections = [
        ("Dept. of CSE,BMSIT&M", 0.5, 2.0, PP_ALIGN.LEFT),
        ("7th Semester, MAJOR PROJECT PHASE-2 (BCS705) REVIEW ‘1", 2.2, 5.6, PP_ALIGN.CENTER),
        ("AY : 2026-27", 7.85, 1.25, PP_ALIGN.RIGHT),
    ]
    for text, left, width, alignment in sections:
        box = slide.shapes.add_textbox(Inches(left), Inches(7.1), Inches(width), Inches(0.28))
        frame = clear_text_frame(box, margin=0)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = alignment
        disable_automatic_bullet(paragraph)
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = BLACK


def set_diagram_title(shape, text, *, size=32):
    frame = clear_text_frame(shape, margin=0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    disable_automatic_bullet(paragraph)
    run = paragraph.add_run()
    run.text = text
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.color.rgb = RED


def move_last_slide_to(presentation, index):
    slide_ids = presentation.slides._sldIdLst
    slide_id = slide_ids[-1]
    slide_ids.remove(slide_id)
    slide_ids.insert(index, slide_id)


def add_diagram_box(slide, title, subtitle, left, top, width, height, *, accent=False, font_size=13):
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    box.fill.solid()
    box.fill.fore_color.rgb = LIGHT_RED if accent else RGBColor(255, 255, 255)
    box.line.color.rgb = RED if accent else MID_GRAY
    box.line.width = Pt(1.15)
    frame = clear_text_frame(box, margin=0.04)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    title_paragraph = frame.paragraphs[0]
    title_paragraph.alignment = PP_ALIGN.CENTER
    disable_automatic_bullet(title_paragraph)
    title_run = title_paragraph.add_run()
    title_run.text = title
    title_run.font.name = BODY_FONT
    title_run.font.size = Pt(font_size)
    title_run.font.bold = True
    title_run.font.color.rgb = BLACK
    if subtitle:
        subtitle_paragraph = frame.add_paragraph()
        subtitle_paragraph.alignment = PP_ALIGN.CENTER
        disable_automatic_bullet(subtitle_paragraph)
        subtitle_paragraph.space_before = Pt(1)
        subtitle_run = subtitle_paragraph.add_run()
        subtitle_run.text = subtitle
        subtitle_run.font.name = BODY_FONT
        subtitle_run.font.size = Pt(max(font_size - 2.5, 9))
        subtitle_run.font.color.rgb = GRAY
    return box


def add_down_arrow(slide, left, top, width=0.28, height=0.28):
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RED
    arrow.line.fill.background()
    return arrow


def add_horizontal_arrow(slide, left, top, *, direction="right", width=0.24, height=0.22):
    arrow_shape = MSO_SHAPE.RIGHT_ARROW if direction == "right" else MSO_SHAPE.LEFT_ARROW
    arrow = slide.shapes.add_shape(
        arrow_shape,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    arrow.fill.solid()
    arrow.fill.fore_color.rgb = RED
    arrow.line.fill.background()
    return arrow


def add_section_label(slide, text, left, top, width):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.28))
    frame = clear_text_frame(box, margin=0)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    disable_automatic_bullet(paragraph)
    run = paragraph.add_run()
    run.text = text.upper()
    run.font.name = BODY_FONT
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = RED
    return box


def add_objective_card(slide, number, title, description, left, top, width, height=1.0):
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(255, 255, 255)
    card.line.color.rgb = MID_GRAY
    card.line.width = Pt(0.85)

    accent = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left + 0.14),
        Inches(top + 0.18),
        Inches(0.48),
        Inches(0.31),
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = LIGHT_RED
    accent.line.fill.background()
    accent_frame = clear_text_frame(accent, margin=0)
    accent_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    accent_paragraph = accent_frame.paragraphs[0]
    accent_paragraph.alignment = PP_ALIGN.CENTER
    disable_automatic_bullet(accent_paragraph)
    accent_run = accent_paragraph.add_run()
    accent_run.text = f"{number:02d}"
    accent_run.font.name = BODY_FONT
    accent_run.font.size = Pt(10.5)
    accent_run.font.bold = True
    accent_run.font.color.rgb = RED

    title_box = slide.shapes.add_textbox(
        Inches(left + 0.75),
        Inches(top + 0.12),
        Inches(width - 0.91),
        Inches(0.3),
    )
    title_frame = clear_text_frame(title_box, margin=0)
    title_paragraph = title_frame.paragraphs[0]
    disable_automatic_bullet(title_paragraph)
    title_run = title_paragraph.add_run()
    title_run.text = title
    title_run.font.name = BODY_FONT
    title_run.font.size = Pt(14)
    title_run.font.bold = True
    title_run.font.color.rgb = BLACK

    description_box = slide.shapes.add_textbox(
        Inches(left + 0.75),
        Inches(top + 0.47),
        Inches(width - 0.91),
        Inches(height - 0.55),
    )
    description_frame = clear_text_frame(description_box, margin=0)
    description_paragraph = description_frame.paragraphs[0]
    disable_automatic_bullet(description_paragraph)
    description_paragraph.line_spacing = 1.0
    description_run = description_paragraph.add_run()
    description_run.text = description
    description_run.font.name = BODY_FONT
    description_run.font.size = Pt(11.5)
    description_run.font.color.rgb = GRAY
    return card


def fill_objectives(slide):
    clear_text_frame(slide.shapes[1], margin=0)
    objectives = [
        ("Architecture First", "Create ARCHITECTURE.md before code generation."),
        ("Approval Gate", "Generate source only after the user confirms the plan."),
        ("Full-Stack Output", "Build a responsive frontend and contract-matched Python backend."),
        ("Prompt-Led Design", "Respect the product name, colors, content, domain and layout."),
        ("Project Workspace", "Provide files, preview, history, versions, restore and export."),
        ("Reliable AI Routing", "Use Gemini, Groq and OpenRouter through controlled fallback."),
        ("Quality & Traceability", "Validate behavior and map requirements to files and tests."),
    ]
    positions = [
        (0.55, 1.60, 4.3),
        (5.15, 1.60, 4.3),
        (0.55, 2.74, 4.3),
        (5.15, 2.74, 4.3),
        (0.55, 3.88, 4.3),
        (5.15, 3.88, 4.3),
        (0.55, 5.02, 8.9),
    ]
    for number, ((title, description), (left, top, width)) in enumerate(zip(objectives, positions), start=1):
        add_objective_card(slide, number, title, description, left, top, width)

    note = slide.shapes.add_textbox(Inches(0.55), Inches(6.20), Inches(8.9), Inches(0.28))
    note_frame = clear_text_frame(note, margin=0)
    note_paragraph = note_frame.paragraphs[0]
    note_paragraph.alignment = PP_ALIGN.CENTER
    disable_automatic_bullet(note_paragraph)
    note_run = note_paragraph.add_run()
    note_run.text = "The goal is a controlled journey from product idea to a working, reviewable application."
    note_run.font.name = BODY_FONT
    note_run.font.size = Pt(11.5)
    note_run.font.italic = True
    note_run.font.color.rgb = GRAY


def add_process_card(slide, number, title, description, left, top, *, final=False):
    width = 2.05
    height = 1.38
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    card.fill.solid()
    card.fill.fore_color.rgb = LIGHT_RED if final else RGBColor(255, 255, 255)
    card.line.color.rgb = RED if final else MID_GRAY
    card.line.width = Pt(1.0)

    badge = slide.shapes.add_textbox(Inches(left + 0.14), Inches(top + 0.12), Inches(0.37), Inches(0.25))
    badge_frame = clear_text_frame(badge, margin=0)
    badge_paragraph = badge_frame.paragraphs[0]
    disable_automatic_bullet(badge_paragraph)
    badge_run = badge_paragraph.add_run()
    badge_run.text = number
    badge_run.font.name = BODY_FONT
    badge_run.font.size = Pt(10.5)
    badge_run.font.bold = True
    badge_run.font.color.rgb = RED

    title_box = slide.shapes.add_textbox(Inches(left + 0.14), Inches(top + 0.39), Inches(1.77), Inches(0.38))
    title_frame = clear_text_frame(title_box, margin=0)
    title_paragraph = title_frame.paragraphs[0]
    disable_automatic_bullet(title_paragraph)
    title_run = title_paragraph.add_run()
    title_run.text = title
    title_run.font.name = BODY_FONT
    title_run.font.size = Pt(13)
    title_run.font.bold = True
    title_run.font.color.rgb = BLACK

    description_box = slide.shapes.add_textbox(Inches(left + 0.14), Inches(top + 0.81), Inches(1.77), Inches(0.43))
    description_frame = clear_text_frame(description_box, margin=0)
    description_paragraph = description_frame.paragraphs[0]
    disable_automatic_bullet(description_paragraph)
    description_paragraph.line_spacing = 1.0
    description_run = description_paragraph.add_run()
    description_run.text = description
    description_run.font.name = BODY_FONT
    description_run.font.size = Pt(10.5)
    description_run.font.color.rgb = GRAY
    return card


def fill_methodology(slide):
    clear_text_frame(slide.shapes[1], margin=0)
    add_section_label(slide, "Planning & approval", 0.55, 1.58, 3.0)
    top_lefts = [0.55, 2.85, 5.15, 7.45]
    top_stages = [
        ("01", "Understand Prompt", "Identify users, features and visual rules."),
        ("02", "Plan Architecture", "Prepare scope, data, APIs and the test plan."),
        ("03", "User Approval", "Edit the prompt or confirm the architecture."),
        ("04", "Generate Full Stack", "Create the React frontend and Python backend."),
    ]
    for index, (left, stage) in enumerate(zip(top_lefts, top_stages)):
        add_process_card(slide, *stage, left, 1.90)
        if index < len(top_lefts) - 1:
            add_horizontal_arrow(slide, left + 2.08, 2.48, direction="right")

    add_down_arrow(slide, 8.33, 3.34, 0.28, 0.38)
    add_section_label(slide, "Build & verify", 0.55, 4.02, 3.0)

    bottom_stages = [
        (7.45, "05", "Engineering Discipline", "Use scoped work, review and bounded repair."),
        (5.15, "06", "Validate Result", "Check build, APIs, controls and accessibility."),
        (2.85, "07", "Version & Trace", "Save files, versions and the review graph."),
        (0.55, "✓", "Ready Project", "Deliver a reviewable full-stack workspace."),
    ]
    for index, (left, number, title, description) in enumerate(bottom_stages):
        add_process_card(slide, number, title, description, left, 4.4, final=number == "✓")
        if index < len(bottom_stages) - 1:
            add_horizontal_arrow(slide, left - 0.27, 4.98, direction="left")

    note = slide.shapes.add_textbox(Inches(1.25), Inches(6.06), Inches(7.5), Inches(0.38))
    note_frame = clear_text_frame(note, margin=0)
    note_paragraph = note_frame.paragraphs[0]
    note_paragraph.alignment = PP_ALIGN.CENTER
    disable_automatic_bullet(note_paragraph)
    note_run = note_paragraph.add_run()
    note_run.text = "Code generation begins only after architecture approval."
    note_run.font.name = BODY_FONT
    note_run.font.size = Pt(12)
    note_run.font.bold = True
    note_run.font.color.rgb = RED


def add_status_panel(slide, heading, items, left, top, width, height, *, accent=False):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    panel.fill.solid()
    panel.fill.fore_color.rgb = LIGHT_RED if accent else LIGHT_GRAY
    panel.line.color.rgb = RED if accent else MID_GRAY
    panel.line.width = Pt(1.0)
    frame = clear_text_frame(panel, margin=0.18)
    heading_paragraph = frame.paragraphs[0]
    disable_automatic_bullet(heading_paragraph)
    heading_paragraph.space_after = Pt(7)
    heading_run = heading_paragraph.add_run()
    heading_run.text = heading
    heading_run.font.name = BODY_FONT
    heading_run.font.size = Pt(16)
    heading_run.font.bold = True
    heading_run.font.color.rgb = RED if accent else BLACK
    for item in items:
        paragraph = frame.add_paragraph()
        disable_automatic_bullet(paragraph)
        paragraph.space_after = Pt(5)
        paragraph.line_spacing = 1.0
        run = paragraph.add_run()
        run.text = f"• {item}"
        run.font.name = BODY_FONT
        run.font.size = Pt(11.8)
        run.font.color.rgb = GRAY
    return panel


def fill_project_status(slide):
    clear_text_frame(slide.shapes[1], margin=0)
    add_status_panel(
        slide,
        "Completed",
        [
            "Architecture-first workflow with a mandatory approval gate.",
            "Prompt-specific frontend, backend and controlled AI fallback.",
            "Files, preview, history, versions, restore, delete and export.",
            "Browser, API and workflow checks for reliable generation.",
        ],
        0.55,
        1.55,
        5.35,
        4.15,
    )
    add_status_panel(
        slide,
        "In Progress",
        ["Integrating an OpenAI API key and a stronger code-generation model for more complex projects."],
        6.15,
        1.55,
        3.3,
        1.88,
        accent=True,
    )
    add_status_panel(
        slide,
        "Next Phase",
        ["Hosted deployment, external Git synchronization and production database adapters."],
        6.15,
        3.67,
        3.3,
        2.03,
    )

    conclusion = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.55),
        Inches(5.96),
        Inches(8.9),
        Inches(0.64),
    )
    conclusion.fill.solid()
    conclusion.fill.fore_color.rgb = RED
    conclusion.line.fill.background()
    conclusion_frame = clear_text_frame(conclusion, margin=0.08)
    conclusion_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    conclusion_paragraph = conclusion_frame.paragraphs[0]
    conclusion_paragraph.alignment = PP_ALIGN.CENTER
    disable_automatic_bullet(conclusion_paragraph)
    conclusion_run = conclusion_paragraph.add_run()
    conclusion_run.text = "ForgeWeb turns a product idea into a working, reviewable full-stack project."
    conclusion_run.font.name = BODY_FONT
    conclusion_run.font.size = Pt(14)
    conclusion_run.font.bold = True
    conclusion_run.font.color.rgb = RGBColor(255, 255, 255)


def add_architecture_slide(presentation, source_slide):
    slide = presentation.slides.add_slide(source_slide.slide_layout)
    set_diagram_title(slide.shapes.title, "SYSTEM ARCHITECTURE")
    for shape in slide.shapes:
        if shape.is_placeholder and int(shape.placeholder_format.type) == 2:
            shape.text = ""
    copy_footer_placeholders(source_slide, slide)

    layers = [
        (
            "USER LAYER",
            1.5,
            0.78,
            [
                ("User Prompt", "Product idea and visual needs"),
                ("ForgeWeb React UI", "Prompt, approval, files and preview"),
            ],
        ),
        (
            "CONTROL PLANE",
            2.56,
            1.05,
            [
                ("Workflow API", "State and confirmation gate"),
                ("Requirement & Architecture Planner", "Product intent and ARCHITECTURE.md"),
                ("Provider Router", "Gemini, Groq and OpenRouter"),
                ("Project Workspace API", "Files, history and export"),
            ],
        ),
        (
            "GENERATION & REVIEW",
            3.94,
            1.12,
            [
                ("Frontend Generator", "React and TypeScript"),
                ("Backend Generator", "Python API and contracts"),
                ("Review & Validation", "Build and interaction checks"),
                ("Code Review Graph", "Requirement-to-file traceability"),
            ],
        ),
        (
            "PROJECT DATA",
            5.42,
            0.92,
            [
                ("Atomic Project Store", "Build and project state"),
                ("Generated Source & Preview", "Frontend, backend and tests"),
                ("Versions & ZIP Export", "Restore and delivery"),
            ],
        ),
    ]

    for label, top, height, modules in layers:
        outer = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(0.45),
            Inches(top),
            Inches(9.08),
            Inches(height),
        )
        outer.fill.solid()
        outer.fill.fore_color.rgb = LIGHT_GRAY
        outer.line.color.rgb = MID_GRAY
        outer.line.width = Pt(0.8)

        label_box = slide.shapes.add_textbox(Inches(0.62), Inches(top + 0.08), Inches(1.15), Inches(height - 0.16))
        label_frame = clear_text_frame(label_box, margin=0)
        label_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        paragraph = label_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.CENTER
        disable_automatic_bullet(paragraph)
        run = paragraph.add_run()
        run.text = label
        run.font.name = BODY_FONT
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RED

        module_left = 1.92
        available_width = 7.38
        gap = 0.13
        module_width = (available_width - gap * (len(modules) - 1)) / len(modules)
        for module_index, (title, subtitle) in enumerate(modules):
            add_diagram_box(
                slide,
                title,
                subtitle,
                module_left + module_index * (module_width + gap),
                top + 0.11,
                module_width,
                height - 0.22,
                accent=module_index == 0,
                font_size=11.5 if len(modules) == 4 else 12,
            )

    add_down_arrow(slide, 4.86, 2.31, 0.28, 0.22)
    add_down_arrow(slide, 4.86, 3.68, 0.28, 0.22)
    add_down_arrow(slide, 4.86, 5.11, 0.28, 0.27)

    note = slide.shapes.add_textbox(Inches(1.0), Inches(6.48), Inches(8.0), Inches(0.3))
    note_frame = clear_text_frame(note, margin=0)
    note_paragraph = note_frame.paragraphs[0]
    note_paragraph.alignment = PP_ALIGN.CENTER
    disable_automatic_bullet(note_paragraph)
    note_run = note_paragraph.add_run()
    note_run.text = "API keys remain in the server layer; generated projects receive only approved source and configuration examples."
    note_run.font.name = BODY_FONT
    note_run.font.size = Pt(11)
    note_run.font.italic = True
    note_run.font.color.rgb = GRAY
    bring_footer_placeholders_to_front(slide)
    add_visible_template_footer(slide)
    return slide


def add_data_flow_slide(presentation, source_slide):
    slide = presentation.slides.add_slide(source_slide.slide_layout)
    set_diagram_title(slide.shapes.title, "DATA FLOW DIAGRAM")
    for shape in slide.shapes:
        if shape.is_placeholder and int(shape.placeholder_format.type) == 2:
            shape.text = ""
    copy_footer_placeholders(source_slide, slide)

    center_left = 2.15
    box_width = 5.7
    box_height = 0.55
    stages = [
        ("User Prompt", "Product features, users, content and visual direction"),
        ("Requirement & Product-Intent Extraction", "Creates structured scope and assumptions"),
        ("ARCHITECTURE.md Proposal", "Frontend, backend, data, security and test plan"),
        ("User Review", "Edit the prompt or confirm the architecture"),
        ("Frontend and Backend Generation", "React interface, Python API and shared contract"),
        ("Review, Tests and Code Graph", "Checks behavior and maps requirements to files"),
        ("Project Workspace", "Files, preview, versions, restore and ZIP export"),
    ]
    top = 1.42
    gap = 0.19
    for index, (title, subtitle) in enumerate(stages):
        current_top = top + index * (box_height + gap)
        if index == 4:
            add_diagram_box(slide, "Frontend Generation", "React and responsive UI", 2.15, current_top, 2.73, box_height, accent=True, font_size=12)
            add_diagram_box(slide, "Backend Generation", "Python API and shared contract", 5.12, current_top, 2.73, box_height, accent=True, font_size=12)
        else:
            add_diagram_box(
                slide,
                title,
                subtitle,
                center_left,
                current_top,
                box_width,
                box_height,
                accent=index in {0, 3, 6},
                font_size=12.5,
            )
        if index < len(stages) - 1:
            add_down_arrow(slide, 4.86, current_top + box_height + 0.015, 0.28, gap - 0.025)

    edit_note = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(0.46),
        Inches(3.62),
        Inches(1.38),
        Inches(0.72),
    )
    edit_note.fill.solid()
    edit_note.fill.fore_color.rgb = LIGHT_RED
    edit_note.line.color.rgb = RED
    edit_frame = clear_text_frame(edit_note, margin=0.04)
    edit_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    edit_paragraph = edit_frame.paragraphs[0]
    edit_paragraph.alignment = PP_ALIGN.CENTER
    disable_automatic_bullet(edit_paragraph)
    edit_run = edit_paragraph.add_run()
    edit_run.text = "Edit returns to the prompt before any source is generated"
    edit_run.font.name = BODY_FONT
    edit_run.font.size = Pt(10.5)
    edit_run.font.color.rgb = RED
    bring_footer_placeholders_to_front(slide)
    add_visible_template_footer(slide)
    return slide


def fill_title_slide(slide):
    frame = clear_text_frame(slide.shapes[0], margin=0.04)

    def line(text="", *, size=16, bold=False, color=BLACK, after=0):
        paragraph = frame.paragraphs[0] if len(frame.paragraphs) == 1 and not frame.paragraphs[0].text else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.CENTER
        paragraph.space_after = Pt(after)
        run = paragraph.add_run()
        run.text = text
        run.font.name = TITLE_FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    line("7th Semester, Major Project Phase-2 (BCS705)", size=20, bold=True, color=PURPLE)
    line("Review – 1 Presentation on", size=18, bold=True, color=PURPLE, after=3)
    line("FORGEWEB", size=25, bold=True)
    line("Architecture-First Full-Stack Web Application Generator", size=18, bold=True, after=3)
    line("[Student Name 1] ([USN 1])", size=14, bold=True)
    line("[Student Name 2] ([USN 2])", size=14, bold=True)
    line("[Student Name 3] ([USN 3])", size=14, bold=True)
    line("[Student Name 4] ([USN 4])", size=14, bold=True, after=3)
    line("VII SEMESTER DEGREE OF BACHELOR OF ENGINEERING IN", size=13, bold=True, color=BLUE)
    line("COMPUTER SCIENCE AND ENGINEERING", size=14, bold=True, color=PURPLE, after=3)
    line("UNDER THE GUIDANCE OF", size=13, bold=True)
    line("[Guide Name]", size=15, bold=True)
    line("[Designation] • DEPT. OF CSE • BMSIT&M", size=13, bold=True, color=PURPLE)


def fill_abstract(slide):
    frame = clear_text_frame(slide.shapes[1], margin=0.05)
    paragraphs = [
        ("Context: ", True, "AI website builders can create pages quickly, but they may repeat layouts, miss requirements or generate code before the user approves the plan."),
        ("Problem: ", True, "A polished screen is not enough when buttons, backend routes, files or security rules are incomplete."),
        ("Objective: ", True, "ForgeWeb turns a product idea into a reviewed full-stack project. It creates requirements and an architecture file, asks for confirmation, and then generates the frontend and backend."),
        ("Method: ", True, "The platform uses React and TypeScript, Gemini for planning and frontend generation, and Groq with OpenRouter fallback for backend generation. It also maintains versions, validation checks and requirement-to-file traceability."),
        ("Result: ", True, "The prototype creates visibly different interfaces for portals, stores, restaurants, inventory tools, schedulers and general websites. Files are inspectable, previews are responsive and project versions can be restored or deleted."),
        ("Keywords: ", True, "Full-stack generation, architecture-first workflow, React, TypeScript, FastAPI, LLM routing, code traceability, responsive UI."),
    ]
    for label, bold, text in paragraphs:
        paragraph = frame.paragraphs[0] if len(frame.paragraphs) == 1 and not frame.paragraphs[0].text else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        disable_automatic_bullet(paragraph)
        paragraph.space_after = Pt(4)
        paragraph.line_spacing = 1.0
        label_run = paragraph.add_run()
        label_run.text = label
        label_run.font.name = BODY_FONT
        label_run.font.size = Pt(15)
        label_run.font.bold = bold
        label_run.font.color.rgb = BLACK
        text_run = paragraph.add_run()
        text_run.text = text
        text_run.font.name = BODY_FONT
        text_run.font.size = Pt(15)
        text_run.font.bold = False
        text_run.font.color.rgb = BLACK


def fill_demonstration(slide, root):
    set_diagram_title(slide.shapes[0], "DEMONSTRATION OF PROJECT EXECUTION", size=25)
    body = slide.shapes[1]
    body.left, body.top, body.width, body.height = Inches(0.5), Inches(1.62), Inches(3.0), Inches(4.95)
    frame = clear_text_frame(body, margin=0.04)
    add_paragraph(frame, "Working flow", size=16, bold=True, space_after=4)
    for text in [
        "Enter a product idea in the prompt box.",
        "Review requirements and ARCHITECTURE.md.",
        "Edit the prompt or confirm generation.",
        "Open frontend, backend, contract, and test files.",
        "Run the responsive preview and save a version.",
    ]:
        add_bullet(frame, text, size=14, space_after=3)
    add_paragraph(frame, "Implemented stack", size=16, bold=True, space_after=3)
    add_paragraph(frame, "React 19 • TypeScript • Vite • Node.js • FastAPI output • Playwright • GSAP • Anime.js • Three.js", size=13, space_after=2)

    add_cropped_picture(slide, root / "forgeweb-preview-desktop.png", 3.72, 1.68, 5.72, 2.15)
    add_caption(slide, "ForgeWeb prompt and architecture-first entry point", 3.72, 3.86, 5.72)
    add_cropped_picture(slide, root / "generated-aura-cream-preview.png", 3.72, 4.16, 2.77, 1.65)
    add_caption(slide, "Prompt-specific client portal", 3.72, 5.84, 2.77)
    add_cropped_picture(slide, root / "generated-ecommerce-preview.png", 6.67, 4.16, 2.77, 1.65)
    add_caption(slide, "Prompt-specific e-commerce site", 6.67, 5.84, 2.77)


def fill_results(slide, root):
    set_diagram_title(slide.shapes[0], "EXPERIMENTAL SETUP, RESULTS & ANALYSIS", size=25)
    body = slide.shapes[1]
    body.left, body.top, body.width, body.height = Inches(0.5), Inches(1.62), Inches(3.36), Inches(4.95)
    frame = clear_text_frame(body, margin=0.04)
    add_paragraph(frame, "Experimental setup", size=15, bold=True, space_after=3)
    add_paragraph(
        frame,
        "The platform was tested with prompts for portals, online stores, restaurants, inventory systems and schedulers. Each generated site was checked in the browser for layout, responsiveness and working controls.",
        size=12.5,
        space_after=7,
        align=PP_ALIGN.LEFT,
    )
    add_paragraph(frame, "Results", size=15, bold=True, space_after=3)
    add_paragraph(
        frame,
        "The websites followed the requested product type, content and visual direction. Navigation, forms, search, file viewing, preview and project actions worked as intended.",
        size=12.5,
        space_after=7,
        align=PP_ALIGN.LEFT,
    )
    add_paragraph(frame, "Analysis", size=15, bold=True, space_after=3)
    add_paragraph(
        frame,
        "Different prompts created visibly different applications instead of one repeated interface. Provider fallback also kept generation available when a model could not respond.",
        size=12.5,
        align=PP_ALIGN.LEFT,
    )

    placements = [
        (root / "generated-project-preview.png", 4.08, 1.68, "Inventory system"),
        (root / "generated-aura-cream-preview.png", 6.82, 1.68, "Client portal"),
        (root / "generated-ecommerce-preview.png", 4.08, 3.78, "E-commerce platform"),
        (root / "generated-restaurant-preview.png", 6.82, 3.78, "Restaurant booking"),
    ]
    for path, left, top, caption in placements:
        add_cropped_picture(slide, path, left, top, 2.48, 1.65)
        add_caption(slide, caption, left, top + 1.68, 2.48)


def main():
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    root = source.parent
    presentation = Presentation(source)
    if len(presentation.slides) != 11:
        raise RuntimeError(f"Expected the unchanged 11-slide template, found {len(presentation.slides)} slides")

    fill_title_slide(presentation.slides[0])

    set_body(
        presentation.slides[1],
        [
            "Abstract",
            "Introduction",
            "Problem Statement",
            "Objectives of the Work",
            "Methodology",
            "System Architecture Diagram",
            "Data Flow Diagram",
            "Demonstration of Project Execution",
            "Experimental Setup, Results and Analysis",
            "Project Status and Conclusion",
            "Publication Status and References",
        ],
        size=22,
    )

    fill_abstract(presentation.slides[2])

    set_body(
        presentation.slides[3],
        [
            "AI can generate code quickly, but fast output is not always complete or dependable.",
            "ForgeWeb converts a plain-language idea into requirements, architecture, code, tests and a live preview.",
            "The user stays in control through a mandatory architecture approval step.",
            "The platform is designed for students, founders and developers who need a working starting point, not only a mock-up.",
            "Its main difference is prompt-specific generation: a restaurant, store or client portal receives its own structure and visual style.",
        ],
        size=20,
    )

    set_paragraph_body(
        presentation.slides[4],
        [
            "People can describe a software idea in simple words, but converting that idea into a complete full-stack application is still difficult. Many AI development tools move directly from the prompt to code without first confirming the real requirements. They may reuse the same design, create disconnected frontend and backend files, or leave important buttons and workflows incomplete. The result can look ready during a demonstration even though the application does not fully match the user’s need.",
            "The main problem is the absence of a guided and trustworthy generation process. A user needs a system that understands the product, prepares a clear architecture for review, generates a prompt-specific frontend and backend only after approval, and then checks whether the complete application works as one connected system. The user should also be able to inspect every file, view the application, track versions and continue the work safely when an AI provider is unavailable.",
        ],
        size=19,
    )

    fill_objectives(presentation.slides[5])
    fill_methodology(presentation.slides[6])

    fill_demonstration(presentation.slides[7], root)
    fill_results(presentation.slides[8], root)

    fill_project_status(presentation.slides[9])

    set_diagram_title(presentation.slides[10].shapes[0], "PUBLICATIONS & REFERENCES", size=27)
    presentation.slides[10].shapes[2].width = Inches(8.79)
    frame = clear_text_frame(presentation.slides[10].shapes[1], margin=0.04)
    add_paragraph(frame, "Publication status", size=17, bold=True, space_after=4)
    add_bullet(frame, "Paper draft in preparation: ‘ForgeWeb: Architecture-First Prompt-to-Full-Stack Generation with Traceability.’", size=15, space_after=4)
    add_bullet(frame, "Target venue will be selected after the complete Phase-2 evaluation; the paper has not yet been submitted.", size=15, space_after=7)
    add_paragraph(frame, "References (IEEE format)", size=17, bold=True, space_after=4)
    references = [
        "[1] T. Thuy and contributors, ‘Code Review Graph,’ GitHub repository, 2025. [Online]. Available: https://github.com/tirth8205/code-review-graph",
        "[2] G. Tan, ‘gstack,’ GitHub repository, 2026. [Online]. Available: https://github.com/garrytan/gstack",
        "[3] multica-ai, ‘andrej-karpathy-skills,’ GitHub repository, 2026. [Online]. Available: https://github.com/multica-ai/andrej-karpathy-skills",
        "[4] Meta Platforms, Inc., ‘React Documentation,’ 2026. [Online]. Available: https://react.dev",
        "[5] S. Ramírez, ‘FastAPI Documentation,’ 2026. [Online]. Available: https://fastapi.tiangolo.com",
        "[6] Google, ‘Gemini API Documentation,’ 2026. [Online]. Available: https://ai.google.dev",
    ]
    for reference in references:
        add_paragraph(frame, reference, size=12.5, space_after=3)

    methodology_slide = presentation.slides[6]
    add_architecture_slide(presentation, methodology_slide)
    move_last_slide_to(presentation, 7)
    add_data_flow_slide(presentation, methodology_slide)
    move_last_slide_to(presentation, 8)

    presentation.save(output)
    print(output)


if __name__ == "__main__":
    main()
