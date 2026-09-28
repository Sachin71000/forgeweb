from __future__ import annotations

import json
import sys
from pathlib import Path

from pptx import Presentation


def font_details(run):
    font = run.font
    return {
        "name": font.name,
        "size_pt": round(font.size.pt, 2) if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "color": (
            str(font.color.rgb)
            if font.color and font.color.type is not None and getattr(font.color, "rgb", None)
            else None
        ),
    }


def main() -> None:
    source = Path(sys.argv[1])
    presentation = Presentation(source)
    if "--summary" in sys.argv:
        for slide_number, slide in enumerate(presentation.slides, start=1):
            print(f"\n--- SLIDE {slide_number}: {slide.slide_layout.name} ---")
            for index, shape in enumerate(slide.shapes):
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    text = shape.text.replace("\x0b", " ").strip()
                    print(f"[{index}] {shape.name}: {text}")
        return
    data = {
        "path": str(source),
        "slide_width_in": round(presentation.slide_width / 914400, 3),
        "slide_height_in": round(presentation.slide_height / 914400, 3),
        "slide_count": len(presentation.slides),
        "slides": [],
    }
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_data = {"number": slide_number, "layout": slide.slide_layout.name, "shapes": []}
        for index, shape in enumerate(slide.shapes):
            shape_data = {
                "index": index,
                "name": shape.name,
                "shape_type": str(shape.shape_type),
                "left_in": round(shape.left / 914400, 3),
                "top_in": round(shape.top / 914400, 3),
                "width_in": round(shape.width / 914400, 3),
                "height_in": round(shape.height / 914400, 3),
                "is_placeholder": shape.is_placeholder,
            }
            if shape.is_placeholder:
                shape_data["placeholder_type"] = str(shape.placeholder_format.type)
            if getattr(shape, "has_text_frame", False):
                shape_data["text"] = shape.text
                paragraphs = []
                for paragraph in shape.text_frame.paragraphs:
                    paragraphs.append(
                        {
                            "level": paragraph.level,
                            "alignment": str(paragraph.alignment),
                            "runs": [
                                {"text": run.text, "font": font_details(run)}
                                for run in paragraph.runs
                            ],
                        }
                    )
                shape_data["paragraphs"] = paragraphs
            if shape.shape_type == 13:
                shape_data["image_type"] = shape.image.ext
            if shape.has_table:
                shape_data["table"] = [[cell.text for cell in row.cells] for row in shape.table.rows]
            slide_data["shapes"].append(shape_data)
        data["slides"].append(slide_data)
    print(json.dumps(data, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
